"""Document parsing strategy protocols and format adapters."""

import asyncio
from pathlib import Path
from typing import Protocol, runtime_checkable
from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document


def _normalize_suffix(suffix: str) -> str:
    """Normalize file extension suffix to lowercase with leading dot."""
    cleaned = suffix.strip().lower()
    if not cleaned:
        return ""
    return cleaned if cleaned.startswith(".") else f".{cleaned}"


@runtime_checkable
class DocumentParser(Protocol):
    """Protocol seam defining file parsing into LangChain Documents."""

    def can_parse(self, suffix: str) -> bool:
        """Check whether the parser supports the given file extension."""
        ...

    def parse(self, path: Path | str) -> list[Document]:
        """Synchronously parse a file into a list of Document objects."""
        ...

    async def aparse(self, path: Path | str) -> list[Document]:
        """Asynchronously parse a file into a list of Document objects."""
        ...


class PDFParser:
    """Parser adapter for PDF documents using PyPDFLoader."""

    SUPPORTED_SUFFIXES: frozenset[str] = frozenset({".pdf"})

    def can_parse(self, suffix: str) -> bool:
        return _normalize_suffix(suffix) in self.SUPPORTED_SUFFIXES

    def parse(self, path: Path | str) -> list[Document]:
        loader = PyPDFLoader(str(path))
        return loader.load()

    async def aparse(self, path: Path | str) -> list[Document]:
        return await asyncio.to_thread(self.parse, path)


class TextParser:
    """Parser adapter for Plaintext (.txt) and Markdown (.md) documents using TextLoader."""

    SUPPORTED_SUFFIXES: frozenset[str] = frozenset({".txt", ".md"})

    def can_parse(self, suffix: str) -> bool:
        return _normalize_suffix(suffix) in self.SUPPORTED_SUFFIXES

    def parse(self, path: Path | str) -> list[Document]:
        loader = TextLoader(str(path), encoding="utf-8")
        return loader.load()

    async def aparse(self, path: Path | str) -> list[Document]:
        return await asyncio.to_thread(self.parse, path)


class DocxParser:
    """Parser adapter for Microsoft Word (.docx) documents using python-docx."""

    SUPPORTED_SUFFIXES: frozenset[str] = frozenset({".docx"})

    def can_parse(self, suffix: str) -> bool:
        return _normalize_suffix(suffix) in self.SUPPORTED_SUFFIXES

    def parse(self, path: Path | str) -> list[Document]:
        file_path = str(path)
        doc = DocxDocument(file_path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": file_path})]

    async def aparse(self, path: Path | str) -> list[Document]:
        return await asyncio.to_thread(self.parse, path)


class CompositeDocumentParser:
    """Composite parser matching file extensions against an ordered registry of parsers."""

    def __init__(self, parsers: list[DocumentParser] | None = None) -> None:
        if parsers is None:
            self._parsers: list[DocumentParser] = [
                PDFParser(),
                TextParser(),
                DocxParser(),
            ]
        else:
            self._parsers = list(parsers)

    @property
    def parsers(self) -> list[DocumentParser]:
        """Return the list of registered parsers."""
        return list(self._parsers)

    def register_parser(self, parser: DocumentParser, prepend: bool = False) -> None:
        """Register a new parser strategy into the composite."""
        if prepend:
            self._parsers.insert(0, parser)
        else:
            self._parsers.append(parser)

    def can_parse(self, suffix: str) -> bool:
        """Check if any registered parser can handle the given suffix."""
        return any(p.can_parse(suffix) for p in self._parsers)

    def _resolve_parser(self, path: Path | str) -> DocumentParser:
        """Find the first matching parser for the file extension or raise ValueError."""
        suffix = Path(path).suffix.lower()
        for p in self._parsers:
            if p.can_parse(suffix):
                return p
        raise ValueError(f"Unsupported file type: {suffix}")

    def parse(self, path: Path | str) -> list[Document]:
        """Parse file using the first registered matching parser."""
        parser = self._resolve_parser(path)
        return parser.parse(path)

    async def aparse(self, path: Path | str) -> list[Document]:
        """Asynchronously parse file using the first registered matching parser."""
        parser = self._resolve_parser(path)
        return await parser.aparse(path)
